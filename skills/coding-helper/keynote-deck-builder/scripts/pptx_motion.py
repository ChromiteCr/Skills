#!/usr/bin/env python3
"""给 python-pptx 生成的片子写入动画与转场。

    python3 pptx_motion.py --inspect deck.pptx   列出每张片分几步、每步动了什么、用什么转场
    python3 pptx_motion.py --selftest

python-pptx 没有动画接口，这里直接写 PresentationML 的 <p:timing> 与转场。
XML 的形状照抄 PowerPoint 自己存出来的文件（PowerPoint for Mac 16.112 加上动画后存盘，
2026-09-17），不是只凭规范推的。四种动作，对应 HTML 模板的四个属性：

  HTML 模板            这里        PowerPoint 里的名字     presetClass / presetID
  data-build 出现      show()     浮入（上浮）            entr / 42
  data-build 退场      hide()     淡化                   exit / 10
  data-dim、data-mark  recolor()  字体颜色                emph / 3
  data-morph           morph()    平滑（Morph）           p159:morph，旧版本回退成淡入淡出

改动过的只有两处数值：时长用 HTML 模板的 320ms（PowerPoint 默认浮入 1 秒、字体颜色 2 秒），
上移距离用舞台高度的 16/1080（PowerPoint 默认 10%）。两个值都是本仓库的取值，不是苹果或微软的参数。
Morph 480ms，同样取自 HTML 模板。

Morph 按对象名配对：相邻两片上各有一个名字以 !! 开头且相同的对象，PowerPoint 就把前一个变到后一个
（微软支持文档「Morph transition: Tips and tricks」）。
"""

import sys
import zipfile
from pathlib import Path

try:
    from lxml import etree
    from pptx import Presentation
    from pptx.oxml import oxml_parser
    from pptx.oxml.ns import qn
except ImportError:
    sys.exit("缺 python-pptx。装：pip3 install python-pptx")

BUILD_MS = 320
MORPH_MS = 480
RISE = 16 / 1080

URI = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main",
}

NS = {"p": URI["p"], "a": URI["a"]}

# (presetID, presetClass, presetSubtype)，取自 PowerPoint 存出的文件
PRESET = {"in": ("42", "entr", "0"), "out": ("10", "exit", "0"), "color": ("3", "emph", "2")}
PRESET_NAME = {("entr", "42"): "浮入", ("entr", "10"): "淡入", ("entr", "1"): "出现",
               ("exit", "10"): "淡出", ("emph", "3"): "字体颜色"}


def _xp(el, expr):
    """python-pptx 的元素改写了 xpath()，不收 namespaces 参数；这里绕过去，任何元素都能用。"""
    return etree._Element.xpath(el, expr, namespaces=NS)


def _clark(name):
    prefix, _, local = name.partition(":")
    return f"{{{URI[prefix]}}}{local}"


def E(name, attrs=None, *children, ns=()):
    """建一个元素。ns 里列的前缀声明在这个元素上，mc:Choice 的 Requires 要求前缀在作用域内。"""
    el = oxml_parser.makeelement(_clark(name), nsmap={k: URI[k] for k in ns} or None)
    for key, value in (attrs or {}).items():
        el.set(_clark(key) if ":" in key else key, str(value))
    el.extend(children)
    return el


# ── 行为节点 ─────────────────────────────────────────────
def _starts(delay):
    return E("p:stCondLst", None, E("p:cond", {"delay": delay}))


def _target(spid):
    return E("p:tgtEl", None, E("p:spTgt", {"spid": spid}))


def _attr_list(name):
    node = E("p:attrName")
    node.text = name
    return E("p:attrNameLst", None, node)


def _visibility(spid, value, delay):
    return E("p:set", None,
             E("p:cBhvr", None,
               E("p:cTn", {"id": 0, "dur": 1, "fill": "hold"}, _starts(delay)),
               _target(spid), _attr_list("style.visibility")),
             E("p:to", None, E("p:strVal", {"val": value})))


def _fade(spid, direction):
    return E("p:animEffect", {"transition": direction, "filter": "fade"},
             E("p:cBhvr", None, E("p:cTn", {"id": 0, "dur": BUILD_MS}), _target(spid)))


def _slide(spid, attr, start, end):
    def point(tm, val):
        return E("p:tav", {"tm": tm}, E("p:val", None, E("p:strVal", {"val": val})))
    return E("p:anim", {"calcmode": "lin", "valueType": "num"},
             E("p:cBhvr", None, E("p:cTn", {"id": 0, "dur": BUILD_MS, "fill": "hold"}),
               _target(spid), _attr_list(attr)),
             E("p:tavLst", None, point(0, start), point(100000, end)))


def _color(spid, rgb):
    return E("p:animClr", {"clrSpc": "rgb", "dir": "cw"},
             E("p:cBhvr", {"override": "childStyle"},
               E("p:cTn", {"id": 0, "dur": BUILD_MS, "fill": "hold"}),
               _target(spid), _attr_list("style.color")),
             E("p:to", None, E("a:srgbClr", {"val": rgb})))


def _effect(kind, spid, grp, node, rgb=None):
    preset, cls, sub = PRESET[kind]
    if kind == "in":
        body = [_visibility(spid, "visible", 0), _fade(spid, "in"),
                _slide(spid, "ppt_x", "#ppt_x", "#ppt_x"),
                _slide(spid, "ppt_y", f"#ppt_y+{RISE:.4f}", "#ppt_y")]
    elif kind == "out":
        body = [_fade(spid, "out"), _visibility(spid, "hidden", BUILD_MS - 1)]
    else:
        body = [_color(spid, rgb)]
    return E("p:par", None,
             E("p:cTn", {"id": 0, "presetID": preset, "presetClass": cls, "presetSubtype": sub,
                         "fill": "hold", "grpId": grp, "nodeType": node},
               _starts(0), E("p:childTnLst", None, *body)))


def _click(effects):
    """一次点击：外层等点击，内层同时开始。第一个效果是 clickEffect，其余 withEffect。"""
    inner = E("p:par", None, E("p:cTn", {"id": 0, "fill": "hold"}, _starts(0),
                               E("p:childTnLst", None, *effects)))
    return E("p:par", None, E("p:cTn", {"id": 0, "fill": "hold"}, _starts("indefinite"),
                              E("p:childTnLst", None, inner)))


def _timing(clicks, builds):
    def slide_target(tag, evt):
        return E(tag, None, E("p:cond", {"evt": evt, "delay": 0}, E("p:tgtEl", None, E("p:sldTgt"))))
    seq = E("p:seq", {"concurrent": 1, "nextAc": "seek"},
            E("p:cTn", {"id": 0, "dur": "indefinite", "nodeType": "mainSeq"},
              E("p:childTnLst", None, *clicks)),
            slide_target("p:prevCondLst", "onPrev"),
            slide_target("p:nextCondLst", "onNext"))
    root = E("p:cTn", {"id": 0, "dur": "indefinite", "restart": "never", "nodeType": "tmRoot"},
             E("p:childTnLst", None, seq))
    timing = E("p:timing", None, E("p:tnLst", None, E("p:par", None, root)))
    if builds:
        # PowerPoint 存盘时按 spid、grpId 排序，照做，读回再存出来就一字不差
        builds = sorted(builds, key=lambda b: (int(b[0]), int(b[1])))
        timing.append(E("p:bldLst", None,
                        *[E("p:bldP", {"spid": s, "grpId": g}) for s, g in builds]))
    # id 按文档顺序编号，和 PowerPoint 存出来的一致
    for n, ctn in enumerate(timing.iter(qn("p:cTn")), 1):
        ctn.set("id", str(n))
    return timing


def _is_text(shape):
    el = shape._element
    return el.tag == qn("p:sp") and el.find(qn("p:txBody")) is not None


# ── 对外接口 ─────────────────────────────────────────────
class Motion:
    """一张片的出场顺序。步数从 1 起，同一步里的动作同时发生，一步一次点击。"""

    def __init__(self, slide):
        self.slide = slide
        self.cues = []

    def show(self, shape, step):
        self._cue(step, "in", shape)

    def hide(self, shape, step):
        self._cue(step, "out", shape)

    def recolor(self, shape, step, rgb):
        self._cue(step, "color", shape, str(rgb).lstrip("#").upper())

    def _cue(self, step, kind, shape, rgb=None):
        if int(step) < 1:
            raise ValueError(f"步数从 1 开始，收到 {step}")
        self.cues.append((int(step), kind, shape, rgb))

    @property
    def steps(self):
        return len({c[0] for c in self.cues})

    def write(self):
        """写进片子。中间空着的步会被合并，因为一次点击里不能什么都不发生。"""
        if not self.cues:
            return 0
        sld = self.slide._element
        if sld.find(qn("p:timing")) is not None:
            raise ValueError("这张片已经有 <p:timing>，不叠写")
        next_grp, builds, clicks = {}, [], []
        for step in sorted({c[0] for c in self.cues}):
            effects = []
            for _, kind, shape, rgb in (c for c in self.cues if c[0] == step):
                spid = str(shape.shape_id)
                grp = next_grp.get(spid, 0)
                next_grp[spid] = grp + 1
                node = "withEffect" if effects else "clickEffect"
                effects.append(_effect(kind, spid, grp, node, rgb))
                if _is_text(shape):
                    builds.append((spid, grp))
            clicks.append(_click(effects))
        sld._insert_timing(_timing(clicks, builds))
        return len(clicks)


def morph(shape, key):
    """给要跨片连续的对象起名。相邻两片上同名的 !! 对象会被 Morph 配成一对。"""
    shape.name = "!!" + str(key)


def transition(slide, kind, dur_ms):
    """kind 为 "morph" 或 "fade"。都包在 mc:AlternateContent 里，不认识的版本走淡入淡出。"""
    sld = slide._element
    for old in sld.findall(qn("p:transition")) + sld.findall(_clark("mc:AlternateContent")):
        sld.remove(old)
    if kind == "morph":
        choice = E("mc:Choice", {"Requires": "p159"},
                   E("p:transition", {"p14:dur": dur_ms},
                     E("p159:morph", {"option": "byObject"}), ns=("p14",)),
                   ns=("p159",))
    elif kind == "fade":
        choice = E("mc:Choice", {"Requires": "p14"},
                   E("p:transition", {"p14:dur": dur_ms}, E("p:fade")),
                   ns=("p14",))
    else:
        raise ValueError(f"不认识的转场 {kind}")
    block = E("mc:AlternateContent", None, choice,
              E("mc:Fallback", None, E("p:transition", None, E("p:fade"))), ns=("mc",))
    anchor = sld.find(qn("p:clrMapOvr"))
    if anchor is None:
        anchor = sld.find(qn("p:cSld"))
    anchor.addnext(block)


# ── 检查 ─────────────────────────────────────────────────
def describe(slide):
    """返回 (转场说明, [[每一步的动作说明]])。"""
    sld = slide._element
    names = {}
    for shape in slide.shapes:
        label = shape.name
        if shape.has_text_frame and shape.text_frame.text.strip() and not label.startswith("!!"):
            label = "「" + shape.text_frame.text.strip().replace("\n", " ")[:14] + "」"
        names[str(shape.shape_id)] = label

    trans = "直接切换"
    block = sld.find(_clark("mc:AlternateContent"))
    node = None
    if block is not None:
        choice = block.find(_clark("mc:Choice"))
        node = choice.find(qn("p:transition")) if choice is not None else None
    if node is None:
        node = sld.find(qn("p:transition"))
    if node is not None and len(node):
        kind = node[0].tag.rsplit("}", 1)[-1]
        dur = node.get(_clark("p14:dur"))
        trans = {"morph": "Morph", "fade": "淡入淡出"}.get(kind, kind) + (f" {dur}ms" if dur else "")

    steps = []
    main = _xp(sld, ".//p:cTn[@nodeType='mainSeq']")
    if main:
        for group in _xp(main[0], "./p:childTnLst/p:par"):
            acts = []
            for ctn in _xp(group, ".//p:cTn[@presetClass]"):
                name = PRESET_NAME.get((ctn.get("presetClass"), ctn.get("presetID")),
                                       f"{ctn.get('presetClass')}/{ctn.get('presetID')}")
                spid = _xp(ctn, ".//p:spTgt/@spid")
                rgb = _xp(ctn, ".//p:to/a:srgbClr/@val")
                acts.append(f"{name} {names.get(spid[0], spid[0]) if spid else '?'}"
                            + (f" → #{rgb[0]}" if rgb else ""))
            steps.append(acts)
    return trans, steps


def inspect(path):
    prs = Presentation(str(path))
    total = 0
    for n, slide in enumerate(prs.slides, 1):
        trans, steps = describe(slide)
        total += len(steps)
        print(f"第 {n} 张  进入：{trans}  {len(steps)} 步")
        for k, acts in enumerate(steps, 1):
            print(f"    {k}. " + "；".join(acts))
    print(f"共 {len(prs.slides)} 张，{total} 次点击动画")


def selftest():
    import tempfile
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(blank)
    boxes = []
    for n in range(3):
        box = s1.shapes.add_textbox(Inches(1), Inches(1 + n * 1.5), Inches(8), Inches(1))
        box.text_frame.text = f"第 {n + 1} 行"
        boxes.append(box)
    m = Motion(s1)
    for n, box in enumerate(boxes, 1):
        m.show(box, n)
        if n > 1:
            m.recolor(boxes[n - 2], n, "747477")
    m.hide(boxes[0], 3)
    m.write()

    s2 = prs.slides.add_slide(blank)
    t2 = s2.shapes.add_textbox(Inches(2), Inches(3), Inches(2), Inches(1))
    t2.text_frame.text = "T"
    morph(t2, "T")
    s3 = prs.slides.add_slide(blank)
    t3 = s3.shapes.add_textbox(Inches(8), Inches(2), Inches(2), Inches(1))
    t3.text_frame.text = "T"
    morph(t3, "T")
    transition(s3, "morph", MORPH_MS)
    transition(s2, "fade", BUILD_MS)

    checks = []
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "t.pptx"
        prs.save(str(out))
        again = Presentation(str(out))                      # 能重新打开
        with zipfile.ZipFile(out) as z:
            x1 = etree.fromstring(z.read("ppt/slides/slide1.xml"))
            x3 = etree.fromstring(z.read("ppt/slides/slide3.xml"))
        order = [c.tag.rsplit("}", 1)[-1] for c in x3]
        checks.append(("子元素顺序 cSld → clrMapOvr → AlternateContent",
                       order[:3] == ["cSld", "clrMapOvr", "AlternateContent"]))
        order1 = [c.tag.rsplit("}", 1)[-1] for c in x1]
        checks.append(("timing 在 clrMapOvr 之后", order1 == ["cSld", "clrMapOvr", "timing"]))
        ids = x1.xpath("//p:cTn/@id", namespaces=NS)
        checks.append(("cTn id 从 1 连续且不重复", ids == [str(i) for i in range(1, len(ids) + 1)]))
        ns = NS
        used = {(c.xpath("string(.//p:spTgt/@spid)", namespaces=ns), c.get("grpId"))
                for c in x1.xpath("//p:cTn[@presetClass]", namespaces=ns)}
        listed = {(b.get("spid"), b.get("grpId")) for b in x1.xpath("//p:bldP", namespaces=ns)}
        checks.append(("bldLst 与各效果的 (spid, grpId) 一一对应", used == listed))
        firsts = [g.xpath("string(.//p:cTn[@presetClass][1]/@nodeType)", namespaces=ns)
                  for g in x1.xpath("//p:cTn[@nodeType='mainSeq']/p:childTnLst/p:par", namespaces=ns)]
        checks.append(("每次点击的第一个效果是 clickEffect", firsts == ["clickEffect"] * 3))
        trans, steps = describe(again.slides[0])
        checks.append(("第 1 张三步，第 3 步有淡出", len(steps) == 3 and any("淡出" in a for a in steps[2])))
        checks.append(("第 3 张是 Morph 480ms", describe(again.slides[2])[0] == "Morph 480ms"))
        checks.append(("第 2 张是淡入淡出 320ms", describe(again.slides[1])[0] == "淡入淡出 320ms"))
        checks.append(("Morph 对象名以 !! 开头", again.slides[2].shapes[0].name == "!!T"))
        try:
            m.write()
            checks.append(("同一张片不叠写 timing", False))
        except ValueError:
            checks.append(("同一张片不叠写 timing", True))

    # 审计复现：--help 原来只打出标题那一行，退出 64
    import subprocess
    run = subprocess.run([sys.executable, str(Path(__file__).resolve()), "--help"],
                         capture_output=True, text=True)
    checks.append(("--help 打印用法（含 --inspect），退出 0",
                   run.returncode == 0 and "--inspect deck.pptx" in run.stdout))

    ok = True
    for name, passed in checks:
        ok &= passed
        print(f"  {'pass' if passed else 'FAIL'}  {name}")
    print("自检" + ("通过" if ok else "未通过"))
    return 0 if ok else 1


def main():
    args = sys.argv[1:]
    if args == ["--selftest"]:
        return selftest()
    if len(args) == 2 and args[0] == "--inspect":
        path = Path(args[1])
        if not path.is_file():
            print(f"找不到文件：{path}")
            return 66
        inspect(path)
        return 0
    # 标题加用法两段；--help 是正常请求，退出 0，其余用错的情况照旧退出 64
    print("\n\n".join(__doc__.strip().split("\n\n")[:2]))
    return 0 if args in (["--help"], ["-h"]) else 64


if __name__ == "__main__":
    raise SystemExit(main())
